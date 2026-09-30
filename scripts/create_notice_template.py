import docx
from docxtpl import DocxTemplate

src_path = r"C:\Users\Gaurav\.gemini\antigravity\brain\49dcb98d-aa68-418e-81ec-39b8a774f8cc\.user_uploaded\media_1790774183369.docx"
template_path = r"templates\notice\template.docx"

doc = docx.Document(src_path)

# 1. Paragraph 9: Academic Year and Date
p9 = doc.paragraphs[9]
p9.runs[0].text = "Academic Year {{ academic_year }}\t           Date:{{ notice_date }}"

# 2. Paragraph 14: Body text with tags
p14 = doc.paragraphs[14]
for r in p14.runs:
    r.text = ""

p14.runs[0].text = "All AIML Club Learners are hereby informed that a session on “"
p14.runs[2].text = "{{ session_title }}"
p14.runs[2].bold = True
p14.runs[4].text = "” will be conducted on "
p14.runs[6].text = "{{ event_date }}"
p14.runs[8].text = " from {{ event_time }} in "
p14.runs[9].text = "{{ venue }} by:"

# 3. Speaker paragraph
# If we put a single tag {{ speakers_text }} or RichText, or use a proper docxtpl loop:
p16 = doc.paragraphs[16]
for r in p16.runs:
    r.text = ""
# In docxtpl, a loop across paragraphs is:
# Paragraph 1: {%p for speaker in speakers %}
# Paragraph 2: {{ speaker }}
# Paragraph 3: {%p endfor %}
# OR we can pass a multiline string or RichText object!
p16.runs[0].text = "{{ speakers_block }}"

# Remove Paragraph 17
p17 = doc.paragraphs[17]
p17._element.getparent().remove(p17._element)

doc.save(template_path)
print(f"Template saved to {template_path}")

tpl = DocxTemplate(template_path)
# Let's test with simple render
from docxtpl import RichText

context = {
    "academic_year": "2026-27",
    "notice_date": "25/09/2026",
    "session_title": "GENERATIVE AI & AGENTIC WORKFLOWS",
    "event_date": "28th September 2026",
    "event_time": "2:00 PM - 4:30 PM",
    "venue": "Seminar Hall 2",
    "speakers_block": "Mr. Alex Turner (AI Architect)\nMiss. Priya Sharma (Lead Researcher)"
}
tpl.render(context)
test_output_path = r"templates\notice\test_rendered_notice.docx"
tpl.save(test_output_path)
print("Render succeeded with speakers_block!")
